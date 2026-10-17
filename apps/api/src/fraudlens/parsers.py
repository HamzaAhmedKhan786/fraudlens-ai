import re
from collections.abc import Callable
from dataclasses import dataclass
from email import policy
from email.parser import BytesParser
from io import BytesIO

from fraudlens.config import Settings


@dataclass(frozen=True)
class ParsedEvidence:
    kind: str
    attributes: dict[str, object]


class EvidenceParser:
    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings or Settings(ocr_enabled=False)

    def parse(self, content_type: str, data: bytes) -> ParsedEvidence:
        parsers: dict[str, Callable[[bytes], ParsedEvidence]] = {
            "message/rfc822": self._email,
            "application/pdf": self._pdf,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": self._docx,
            "text/plain": self._text,
            "image/png": self._image,
            "image/jpeg": self._image,
        }
        parser = parsers.get(content_type)
        if not parser:
            raise ValueError(f"No parser for {content_type}")
        return parser(data)

    def _email(self, data: bytes) -> ParsedEvidence:
        message = BytesParser(policy=policy.default).parsebytes(data)
        bodies: list[str] = []
        attachments: list[str] = []
        for part in message.walk():
            disposition = part.get_content_disposition()
            if disposition == "attachment":
                attachments.append(part.get_filename() or "unnamed")
            elif part.get_content_type() == "text/plain" and disposition != "attachment":
                try:
                    bodies.append(part.get_content())
                except (LookupError, UnicodeError):
                    pass
        authentication = "\n".join(
            str(value) for value in message.get_all("Authentication-Results", [])
        )
        received_spf = "\n".join(str(value) for value in message.get_all("Received-SPF", []))
        auth_lower = authentication.lower()
        return ParsedEvidence(
            "email",
            {
                "subject": str(message.get("Subject", ""))[:500],
                "from": str(message.get("From", ""))[:500],
                "to": str(message.get("To", ""))[:500],
                "reply_to": str(message.get("Reply-To", ""))[:500],
                "message_id": str(message.get("Message-ID", ""))[:500],
                "attachments": attachments[:100],
                "text_excerpt": "\n".join(bodies)[:10_000],
                "authentication_results": authentication[:4000],
                "spf_pass": "spf=pass" in auth_lower or "pass" in received_spf.lower(),
                "dkim_pass": "dkim=pass" in auth_lower,
                "dmarc_pass": "dmarc=pass" in auth_lower,
            },
        )

    def _pdf(self, data: bytes) -> ParsedEvidence:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data), strict=False)
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        ocr = (
            self._ocr_pdf(data, text)
            if self.settings.ocr_enabled and len(text.strip()) < 40
            else {
                "ocr_status": "NOT_REQUIRED" if text.strip() else "DISABLED",
                "ocr_text": "",
                "ocr_language": "unknown",
                "ocr_word_confidence": None,
                "ocr_pages": [],
                "ocr_error": None,
            }
        )
        combined_text = text or str(ocr.get("ocr_text", ""))
        metadata = reader.metadata or {}
        return ParsedEvidence(
            "document",
            {
                "page_count": len(reader.pages),
                "title": str(metadata.get("/Title", ""))[:500],
                "author": str(metadata.get("/Author", ""))[:500],
                "producer": str(metadata.get("/Producer", ""))[:500],
                "text_excerpt": combined_text[:10_000],
                "encrypted": bool(reader.is_encrypted),
                **ocr,
                **self._document_signals(combined_text),
            },
        )

    def _docx(self, data: bytes) -> ParsedEvidence:
        from docx import Document

        document = Document(BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        tables = [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        text = "\n".join([*paragraphs, *tables])
        return ParsedEvidence(
            "document",
            {
                "paragraph_count": len(paragraphs),
                "table_row_count": len(tables),
                "text_excerpt": text[:10_000],
                **self._document_signals(text),
            },
        )

    def _text(self, data: bytes) -> ParsedEvidence:
        text = data.decode("utf-8", errors="replace")
        return ParsedEvidence(
            "document",
            {
                "character_count": len(text),
                "text_excerpt": text[:10_000],
                **self._document_signals(text),
            },
        )

    def _image(self, data: bytes) -> ParsedEvidence:
        from PIL import ExifTags, Image

        with Image.open(BytesIO(data)) as image:
            exif = {
                str(ExifTags.TAGS.get(key, key)): str(value)[:500]
                for key, value in image.getexif().items()
            }
            ocr = (
                self._ocr_images([image.copy()])
                if self.settings.ocr_enabled
                else {
                    "ocr_status": "DISABLED",
                    "ocr_text": "",
                    "ocr_language": "unknown",
                    "ocr_word_confidence": None,
                    "ocr_pages": [],
                    "ocr_error": None,
                }
            )
            return ParsedEvidence(
                "image",
                {
                    "width": image.width,
                    "height": image.height,
                    "format": image.format or "unknown",
                    "mode": image.mode,
                    "exif_keys": sorted(exif.keys())[:100],
                    **ocr,
                    **self._document_signals(str(ocr.get("ocr_text", ""))),
                },
            )

    def _ocr_pdf(self, data: bytes, existing_text: str) -> dict[str, object]:
        try:
            from pdf2image import convert_from_bytes

            images = convert_from_bytes(
                data,
                first_page=1,
                last_page=self.settings.ocr_max_pages,
                poppler_path=self.settings.poppler_path,
                dpi=220,
            )
            return self._ocr_images(images, existing_text)
        except Exception as exc:
            return self._ocr_failure(exc)

    def _ocr_images(self, images, hint_text: str = "") -> dict[str, object]:
        try:
            import pytesseract
            from pytesseract import Output

            pytesseract.pytesseract.tesseract_cmd = self.settings.tesseract_cmd
            languages, detected = self._select_ocr_languages(images[0], hint_text)
            config = ""
            if self.settings.tessdata_prefix:
                config = f"--tessdata-dir {self.settings.tessdata_prefix}"
            pages: list[dict[str, object]] = []
            all_text: list[str] = []
            confidences: list[float] = []
            for page_number, image in enumerate(images, start=1):
                result = pytesseract.image_to_data(
                    image, lang=languages, config=config, output_type=Output.DICT
                )
                words: list[str] = []
                page_confidences: list[float] = []
                for word, confidence in zip(result["text"], result["conf"], strict=False):
                    word = str(word).strip()
                    try:
                        numeric_confidence = float(confidence)
                    except (TypeError, ValueError):
                        continue
                    if word and numeric_confidence >= 0:
                        words.append(word)
                        page_confidences.append(numeric_confidence)
                        confidences.append(numeric_confidence)
                page_text = " ".join(words)
                all_text.append(page_text)
                pages.append(
                    {
                        "page": page_number,
                        "text": page_text[:10_000],
                        "word_confidence": round(sum(page_confidences) / len(page_confidences), 2)
                        if page_confidences
                        else None,
                    }
                )
            return {
                "ocr_status": "COMPLETED",
                "ocr_text": "\n".join(all_text)[:50_000],
                "ocr_language": detected,
                "ocr_languages_used": languages.split("+"),
                "ocr_word_confidence": round(sum(confidences) / len(confidences), 2)
                if confidences
                else None,
                "ocr_pages": pages,
                "ocr_error": None,
                "ocr_authenticity_verified": False,
            }
        except Exception as exc:
            return self._ocr_failure(exc)

    def _select_ocr_languages(self, image, hint_text: str) -> tuple[str, str]:
        configured = self.settings.ocr_languages.split("+")
        if re.search(r"[\u0600-\u06ff]", hint_text):
            selected = [language for language in ("ara", "urd") if language in configured]
            return "+".join(selected or configured), "Arabic script"
        try:
            import pytesseract

            osd = pytesseract.image_to_osd(image)
            script_match = re.search(r"Script:\s*([^\n]+)", osd)
            script = script_match.group(1).strip() if script_match else "unknown"
            if script == "Arabic":
                selected = [language for language in ("ara", "urd") if language in configured]
                return "+".join(selected or configured), "Arabic script"
            if script == "Latin":
                selected = [
                    language
                    for language in ("eng", "deu", "spa", "ita", "fra")
                    if language in configured
                ]
                return "+".join(selected or configured), "Latin script"
            return "+".join(configured), script
        except Exception:
            return "+".join(configured), "unknown"

    @staticmethod
    def _ocr_failure(exc: Exception) -> dict[str, object]:
        return {
            "ocr_status": "FAILED",
            "ocr_text": "",
            "ocr_language": "unknown",
            "ocr_word_confidence": None,
            "ocr_pages": [],
            "ocr_error": str(exc)[:500],
            "ocr_authenticity_verified": False,
        }

    @staticmethod
    def _document_signals(text: str) -> dict[str, bool | list[str]]:
        normalized = " ".join(text.lower().split())
        country_lengths = {
            "AT": 20,
            "BE": 16,
            "CH": 21,
            "DE": 22,
            "ES": 24,
            "FR": 27,
            "GB": 22,
            "IE": 22,
            "IT": 27,
            "NL": 18,
            "PL": 28,
        }
        ibans: set[str] = set()
        for match in re.findall(r"\b[A-Z]{2}\d{2}(?:[ ]?[A-Z0-9]){11,30}\b", text, re.IGNORECASE):
            candidate = re.sub(r"\s+", "", match).upper()
            expected_length = country_lengths.get(candidate[:2])
            if expected_length:
                candidate = candidate[:expected_length]
            rearranged = candidate[4:] + candidate[:4]
            numeric = "".join(
                character if character.isdigit() else str(ord(character) - 55)
                for character in rearranged
            )
            if 15 <= len(candidate) <= 34 and int(numeric) % 97 == 1:
                ibans.add(candidate)
        payment_change_phrases = (
            "bank details have changed",
            "new bank account",
            "updated bank account",
            "change of bank details",
            "pay to the following account",
            "use the new iban",
        )
        urgency_phrases = (
            "urgent payment",
            "pay immediately",
            "immediate transfer",
            "payment today",
            "avoid delay",
            "confidential payment",
        )
        return {
            "ibans": sorted(ibans)[:20],
            "payment_change_language": any(
                phrase in normalized for phrase in payment_change_phrases
            ),
            "urgent_payment_language": any(phrase in normalized for phrase in urgency_phrases),
        }
